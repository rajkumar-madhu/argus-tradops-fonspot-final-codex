from logging import exception
from django.shortcuts import render
import json
import requests
from requests.auth import HTTPBasicAuth
from analytics.models import UserSettingsModel
from django.contrib.auth.models import User
from django.http import HttpResponse
from django.http import HttpResponseRedirect
from django.contrib.auth.decorators import login_required
import os
import psycopg2
from django.conf import settings
from login.decorators import role_required
from lib.LinkedEyeVault.AppSecrets import get_app_secret
from elasticsearch import Elasticsearch
from django.http import JsonResponse
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from reportlab.lib.pagesizes import letter, landscape
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle
from datetime import datetime, timedelta
from reportlab.pdfgen import canvas
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import inch
from io import BytesIO

class AnalyticsConnection:
    """Context manager for analytics PostgreSQL connections. HIGH FIX #11."""
    def __enter__(self):
        self.conn = psycopg2.connect(
            database=settings.POSTGRES_SUPERSET_DB,
            user=settings.POSTGRES_USER,
            password=settings.POSTGRES_PASS,
            host=settings.POSTGRES_HOST,
            port=settings.POSTGRES_PORT
        )
        self.cursor = self.conn.cursor()
        return self.cursor

    def __exit__(self, exc_type, exc_val, exc_tb):
        if self.cursor:
            self.cursor.close()
        if self.conn:
            self.conn.close()
        return False

def setup_connection():
    """Legacy wrapper — returns cursor. Caller must close connection.
    Prefer using `with AnalyticsConnection() as cursor:` instead."""
    conn = psycopg2.connect(
        database=settings.POSTGRES_SUPERSET_DB,
        user=settings.POSTGRES_USER,
        password=settings.POSTGRES_PASS,
        host=settings.POSTGRES_HOST,
        port=settings.POSTGRES_PORT
    )
    cursor = conn.cursor()
    cursor._analytics_conn = conn  # Attach connection for cleanup
    return cursor

json_path = "management-console/device-templates/"

@login_required(login_url="/")
#@role_required(allowed_roles=["Admin", "Management", "ViewOnly"])

def getTableIndex(request):
    response  = {}
    response['data'] = {}
    response['data']['tables'] = {}
    try: 
        cursor = setup_connection()
        cursor.execute("select id, table_name from tables;")
        tables = cursor.fetchall()
        for table in tables:
            response['data']['tables'][table[1]] = str(table[0])+'__table'
        response['status'] = 200
    except Exception as ex:
        response['status'] = 500
        response['msg'] = str(ex)
    #print(" getTableIndex ---> {}".format(response))
    return HttpResponse(json.dumps(response), content_type="json")

@role_required(allowed_roles=["Admin", "Management", "ViewOnly", "Onboard", "Risk"])
def dashboard(request):
    """
    jsonname = request.GET["jsonname"]
    tableinfo = request.GET.get("tableinfo", {})
    table_id = json.loads(tableinfo).get('tables', {})
    #jsonname = "noren-oms.json"  
    filedata = open(json_path+jsonname)
    filecontent = filedata.read()
    filedata.close()
    data = json.loads(filecontent) # deserialises it
    totaltable = data['totaltable']
    #----
    for k,v in table_id.items():
        print("{} {}".format(k,v))
        filecontent = filecontent.replace(k,v)
    data = json.loads(filecontent)
    #----
    
    if 'totaltable' not in data:
        temp_json = {"status": 500,
                    "msg": "Analytics Data not able to Fetch. Please contact Administrator",
                    }
    else:
        wrong_tables = []
        for k,v in totaltable.items():
            print("{} {}".format(k, list(table_id.keys())))
            if not k in list(table_id.keys()):
                wrong_tables.append(k)
        if wrong_tables == []:
            data['totaltable'] = totaltable
            json_data = json.dumps(data) # json formatted string      
            if UserSettingsModel.objects.filter(username = request.user,jsonid=data['uid']).exists():
                settings = UserSettingsModel.objects.get(username = request.user,jsonid=data['uid']).analytics
            else:
                settings = { }
            temp_json = {"status":200,
                        "data": json_data,
                        "settings": json.dumps(settings),
                        }
        else:
            tables = ", ".join(str(table) for table in wrong_tables)
            temp_json = {"status": 500,
                    "msg": "Table [" + tables + "] not found. Please contact Administrator",
                    } 

    return render(request, 'app/analytics.html', temp_json)
            """

    return render(request, 'app/analytics.html')

def saveSettings(request):
    if request.method == 'POST':
        response = { }
        try:
            parsed_json = request.POST['settingsData']
            if UserSettingsModel.objects.filter(username = request.user,jsonid=json.loads(parsed_json)['id']).exists():
                obj = UserSettingsModel.objects.get(username = request.user,jsonid=json.loads(parsed_json)['id'])
                obj.analytics = parsed_json
                obj.save()
            else:
                userid = User.objects.get(username = request.user).id
                obj = UserSettingsModel(username = request.user, user_id = userid, analytics = parsed_json,jsonid=json.loads(parsed_json)['id'])
                obj.save()
            response['status'] = 200
            response['msg'] = 'Successfully saved.'
        except Exception as e:
            #print('==Exception====saveSettings=')
            #print(str(e))
            response['status'] = 400
            response['msg'] = 'Not able to save'
    return HttpResponse(json.dumps(response), content_type="json")

def getprefixurlData(request):
    #print(request)
    userData = {
        'user': settings.ANALYTICS_DASHBOARD_USER,
        "analysticsDashboardurl": request.GET.get('url')
    }
    #print('this is analytics views')
    #print(userData)
    json_data = json.dumps(userData)
    #print('this is analytics views json_data')
    #print(json_data)
    return HttpResponse(json_data)


def getpermalink(request):
    response = {}
    try:    
        #print(request)
        response['status'] = 200
        access_token=request.POST['accesstoken']
        #url='http://172.16.0.22:8088/api/v1/explore/permalink'
      #  url='http://172.16.0.22:8088/api/v1/explore/permalink?Bearer='+access_token
        url=request.POST['url']+'api/v1/explore/permalink'
        formdata=json.loads(request.POST['formdata'])
        urlparams=[]
        #urlparams=request.POST['urlparams']
        params = {
            "formData": formdata,
            "urlParams": urlparams
        }
        headers = {"Authorization": "Bearer "+access_token,'Content-Type': 'application/json'}
        #print('this is analytics views')
        #print(" -----> {} {} {}".format(url,headers,params))
        permalink=requests.post(url,data=json.dumps(params),headers=headers)
        if permalink.ok:
            response = permalink.json()
        else:
            #print("Permalink Get Failed")
            #print(permalink.content)
            response['status'] = 400
            response['msg'] = 'Permalink Get Failed'
            response['error'] = permalink.content
    except Exception as e:
            #print('==Exception====GetPermalink=')
            #print(str(e))
            response['status'] = 400
            response['msg'] = 'Not able to Get Permalink with err message : ' + str(e) 
    #print('this is analytics views json_data')
    #print(response)
    return HttpResponse(json.dumps(response, default=str), content_type="json")

"""
def getpermalink(request):
    print(request)
    response = {}
    response['status'] = 200
    access_token=request.POST['accesstoken']
    #url='http://172.16.0.22:8088/api/v1/explore/permalink'
  #  url='http://172.16.0.22:8088/api/v1/explore/permalink?Bearer='+access_token
    url=request.POST['url']+'api/v1/explore/permalink'
    formdata=json.loads(request.POST['formdata'])
    urlparams=[]
    #urlparams=request.POST['urlparams']
    params = {
        "formData": formdata,
        "urlParams": urlparams
    }
    headers = {"Authorization": "Bearer "+access_token,'Content-Type': 'application/json'}
    print('this is analytics views')
    print(" -----> {} {} {}".format(url,headers,params))
    permalink=requests.post(url,data=json.dumps(params),headers=headers)
    if permalink.ok:
        response = permalink.json()
    else:
        print("Permalink Get Failed")
        print(permalink.content)
        response['status'] = 400
        response['msg'] = 'Permalink Get Failed'
        response['error'] = permalink.content
    print('this is analytics views json_data')
    print(response)
    return HttpResponse(json.dumps(response, default=str), content_type="json")
"""

        
def getaccesstoken(request):
    response={}
    try:
        #print(request)
        url=request.POST['url']+'api/v1/security/login'
        #url='http://172.16.0.22:8088/api/v1/security/login'
        # Superset service-account credentials. These were hardcoded here (with
        # the real password in a trailing comment) and committed -- resolved from
        # Vault/env now, matching how every other service credential in this
        # project is handled. The old value is still in git history, so it has to
        # be rotated in Superset itself; removing it here is not sufficient.
        superset_user = get_app_secret('SUPERSET_USER', env_var='SUPERSET_USER', default='')
        superset_password = get_app_secret('SUPERSET_PASSWORD', env_var='SUPERSET_PASSWORD', default='')
        param = {
            "password": superset_password,
            "provider": "db",
            "refresh": "true",
            "username": superset_user
        }
        #print('this is getaccesstoken')
        token_json=requests.post(url = url, json = param,auth=HTTPBasicAuth(superset_user, superset_password))
        #token_json = json.dumps(token_json)
        
        response['url']=request.POST['url']
        response['token_json']=token_json.json()
        #print(response)
    except Exception as e:
            #print('==Exception====GetAccessToken=')
            #print(str(e))
            response['status'] = 400
            response['msg'] = 'Not able to Get AccessToken with err message : ' + str(e) 
    #print(json_data)
    #return HttpResponse(response)
    return HttpResponse(json.dumps(response, default=str), content_type="json")

def getUID(request):
    response = {}
    try:
        from lib.LinkedEyeMonitoring.token import monitoring_credentials, verify_ssl
        url = request.GET['url'] + '/api/search?query=' + request.GET['dbname']

        grafana_auth, grafana_headers = monitoring_credentials('grafana')
        headers = {'Content-Type': 'application/json'}
        headers.update(grafana_headers)

        token_json = requests.get(
            url=url,
            headers=headers,
            auth=grafana_auth,
            verify=verify_ssl()
        )

        response['url'] = request.GET['url']
        response['token_json'] = token_json.json()
        db_uid=response['token_json'][0]['uid']
        url = request.GET['url'] + '/api/dashboards/uid/' + db_uid
        response['db_json'] = requests.get(
            url=url,
            headers=headers,
            auth=grafana_auth,
            verify=verify_ssl()
        ).json()
    except Exception as e:
        #print('==Exception====GetUID=')
        #print(str(e))
        response['status'] = 400
        response['msg'] = 'Not able to Get UID with err message: ' + str(e)

    return HttpResponse(json.dumps(response, default=str), content_type="application/json")

def monitorgraph(request):
    graphdata =  request.POST['service']
    temp_json = {"status":200, "graphdata": graphdata}
    return render(request, 'app/monitor.html', temp_json)



def search_elasticsearch(request):
    try:
        #index_name = 'noren-login-history'
        # print('elastic host--->{}'.format((request.GET.get("elastic_host", "172.20.1.80"))))
        # print('elastic port--->{}'.format(request.GET.get("elastic_port", 31545)))
        # Read subsite and index_name from JS
        subsite = request.GET.get("subsite", "").strip()
        index_name = request.GET.get("index_name", "").strip()

        # Safety fallback
        if not index_name:
            if subsite:
                index_name = f"{subsite}-login-history"
            else:
                index_name = "noren-login-history"

        #print("Subsite received:", subsite)
        #print("Index selected:", index_name)

        es = Elasticsearch(
            [{'host': request.GET.get("elastic_host", settings.ELASTIC_HOST),
                'port': int(request.GET.get("elastic_port", settings.ELASTIC_PORT)),
                'scheme': 'http'
            }],
            http_auth=(settings.ELASTIC_USER, settings.ELASTIC_PASS)
            )


        # print('elastic user--->{}'.format(os.getenv('ELASTIC_USER', 'elastic')))
        # print('elastic pass--->{}'.format(os.getenv('ELASTIC_PASS', 'changeme')))
        # Pagination parameters
        start = int(request.GET.get("start", 0))  # starting record (pagination)
        length = int(request.GET.get("length", 50))  # records per page (50)

        # Parse columns and order from the request
        columns = []
        i = 0
        while f"columns[{i}][data]" in request.GET:
            columns.append({
                "data": request.GET.get(f"columns[{i}][data]"),
                "search": request.GET.get(f"columns[{i}][search]")
            })
            i += 1

        # Parse order
        order = []
        i = 0
        while f"order[{i}][column]" in request.GET:
            order_column = request.GET.get(f"order[{i}][column]")
            order_dir = request.GET.get(f"order[{i}][dir]")

            try:
                column_index = int(order_column)
                column_name = columns[column_index]["data"] if column_index < len(columns) else None
            except ValueError:
                column_name = order_column

            if column_name:
                order.append({
                    "column": column_name,
                    "dir": order_dir
                })
            i += 1

        # Prepare the Elasticsearch query for getting the total number of records matching the query
        count_query = {
            "query": {
                "bool": {
                    "must": []
                }
            }
        }

        # Apply filters (start_time, end_time, etc.) to the count query
        start_time = request.GET.get("start_time", None)
        end_time = request.GET.get("end_time", None)
        if start_time and end_time:
            count_query["query"]["bool"]["must"].append({
                "range": {
                    "@timestamp": {
                        "gte": start_time,
                        "lte": end_time,
                        "format": "strict_date_optional_time"
                    }
                }
            })

        for col in columns:
            filter_value = col["search"]
            if filter_value:
                if col["data"] in ["UserId", "UserName", "BrokerId", "LastLoginIp", "LastLoginMac"]:
                    count_query["query"]["bool"]["must"].append({
                        "wildcard": {
                            f"Userdetails.{col['data']}.keyword": {
                                "value": f"*{filter_value}*",
                                "case_insensitive": True
                            }
                        }
                    })
                elif col["data"] == "@timestamp":
                    count_query["query"]["bool"]["must"].append({
                        "wildcard": {
                            "@timestamp_string": {
                                "value": f"*{filter_value}*",
                                "case_insensitive": True
                            }
                        }
                    })
                else:
                    count_query["query"]["bool"]["must"].append({
                        "query_string": {
                            "fields": [col["data"]],
                            "query": f"*{filter_value.lower()}*",
                            "analyze_wildcard": True,
                            "default_operator": "AND"
                        }
                    })

        # Fetch the total number of matching records without fetching the actual data
        # print('COUNT_QUERY--->{}'.format(count_query))
        count_response = es.count(index=index_name, body=count_query)
        total_records = count_response["count"]  # Get the total count of matching records

        # Prepare the Elasticsearch query to fetch the current page of records (pagination)
        query = {
            "_source": [
                "UserId",
                "Userdetails.UserName",
                "Userdetails.BrokerId",
                "@timestamp",
                "ReqStatus",
                "AccessType",
                "Userdetails.LastLoginIp",
                "Userdetails.LastLoginMac"
            ],
            "size": length,  # only return 50 records per request
            "from": start,  # Pagination (starting point)
            "query": {
                "bool": {
                    "must": []
                }
            },
        }

        # Apply the date range filter to the main query
        if start_time and end_time:
            query["query"]["bool"]["must"].append({
                "range": {
                    "@timestamp": {
                        "gte": start_time,
                        "lte": end_time,
                        "format": "strict_date_optional_time"
                    }
                }
            })

        # Apply sorting if necessary
        if order:
            sort_field = order[0]["column"]
            sort_dir = order[0]["dir"]
            if sort_field in ["UserId", "UserName", "BrokerId", "LastLoginIp", "LastLoginMac"]:
                sort_field = 'Userdetails.' + sort_field + ".keyword"
            elif sort_field == "@timestamp":
                sort_field = "@timestamp"
            else:
                sort_field += ".keyword"
            query["sort"] = [{sort_field: {"order": sort_dir}}]

        # Apply filters for the page query (date range, text search, etc.)
        for col in columns:
            filter_value = col["search"]
            if filter_value:
                if col["data"] in ["UserId", "UserName", "BrokerId", "LastLoginIp", "LastLoginMac"]:
                    query["query"]["bool"]["must"].append({
                        "wildcard": {
                            f"Userdetails.{col['data']}.keyword": {
                                "value": f"*{filter_value}*",
                                "case_insensitive": True
                            }
                        }
                    })
                elif col["data"] == "@timestamp":
                    query["query"]["bool"]["must"].append({
                        "wildcard": {
                            "@timestamp_string": {
                                "value": f"*{filter_value}*",
                                "case_insensitive": True
                            }
                        }
                    })
                else:
                    query["query"]["bool"]["must"].append({
                        "query_string": {
                            "fields": [col["data"]],
                            "query": f"*{filter_value.lower()}*",
                            "analyze_wildcard": True,
                            "default_operator": "AND"
                        }
                    })
        #print('Main_QUERY--->{}'.format(query))
        # Fetch the first batch of results from Elasticsearch (50 records for the page)
        response = es.search(index=index_name, body=query)

        # Collect data from the first batch of results
        all_hits = response['hits']['hits']
        data = [hit["_source"] for hit in all_hits]

        # Format data for DataTable (only the first `length` records)
        def flatten_hit(hit):
            source = hit
            timestamp_utc = source.get("@timestamp", "")
            human_readable_timestamp = ""
            if timestamp_utc:
                try:
                    utc_time = datetime.strptime(timestamp_utc, "%Y-%m-%dT%H:%M:%S.%fZ")
                    ist_time = utc_time + timedelta(hours=5, minutes=30)
                    human_readable_timestamp = ist_time.strftime("%Y-%m-%d %I:%M:%S %p")
                except Exception as e:
                    human_readable_timestamp = timestamp_utc
            return {
                "UserId": source.get("UserId", ""),
                "UserName": source.get("Userdetails", {}).get("UserName", ""),
                "BrokerId": source.get("Userdetails", {}).get("BrokerId", ""),
                "@timestamp": human_readable_timestamp,
                "ReqStatus": source.get("ReqStatus", ""),
                "AccessType": source.get("AccessType", ""),
                "LastLoginIp": source.get("Userdetails", {}).get("LastLoginIp", ""),
                "LastLoginMac": source.get("Userdetails", {}).get("LastLoginMac", ""),
            }

        # Limit data to only the first `length` records for pagination
        formatted_data = [flatten_hit(hit) for hit in data[:length]]  # Only 50 records for the page

        # Return response in DataTables format
        return JsonResponse({
            "draw": int(request.GET.get("draw", 1)),
            "recordsTotal": total_records,  # Total records matching the query
            "recordsFiltered": total_records,  # Filtered records (same as total in this case)
            "status": 200,
            "results": formatted_data
        })

    except Exception as e:
        #print('EXCEPTION - {}'.format(e))
        return JsonResponse({"status": 400, "results": str(e)}, status=400)

def export_to_excel(request):
    if request.method != 'POST':
        return JsonResponse({'error': 'Invalid HTTP method. Use POST.'}, status=405)

    try:
        # Parse JSON body
        data = json.loads(request.POST['req'])  # Ensure correct parsing
        filters = data.get('filters', {})
        sorting = data.get('sorting', [])

        # Configure Elasticsearch
        es = Elasticsearch([{'host': request.GET.get("elastic_host", settings.ELASTIC_HOST), 'port': int(request.GET.get("elastic_port", settings.ELASTIC_PORT)), 'scheme': 'http'}], http_auth=(settings.ELASTIC_USER, settings.ELASTIC_PASS))

        index_name = 'noren-login-history'

        # Build Elasticsearch query
        query = {
            "_source": [
                "UserId",
                "Userdetails.UserName",
                "Userdetails.BrokerId",
                "@timestamp",
                "ReqStatus",
                "AccessType",
                "Userdetails.LastLoginIp",
                "Userdetails.LastLoginMac"
            ],
            "query": {
                "bool": {
                    "must": []
                }
            },
            "sort": []
        }

        # Add filters
        for col, value in filters.items():
            if value:
                if col in ["UserId", "UserName", "BrokerId", "LastLoginIp", "LastLoginMac"]:
                    query["query"]["bool"]["must"].append({
                        "wildcard": {
                            f"Userdetails.{col}.keyword": {
                                "value": f"*{value}*",
                                "case_insensitive": True
                            }
                        }
                    })
                elif col == "@timestamp":
                    query["query"]["bool"]["must"].append({
                        "range": {
                            "@timestamp": {
                                "gte": filters.get('start_time'),
                                "lte": filters.get('end_time'),
                                "format": "strict_date_optional_time"
                            }
                        }
                    })

        # Add sorting
        for sort in sorting:
            column = sort.get('column', '').replace('.', '_')
            direction = sort.get('dir', 'asc')
            if column in ["UserId", "UserName", "BrokerId", "LastLoginIp", "LastLoginMac"]:
                column = f"Userdetails.{column}.keyword"
            query["sort"].append({column: {"order": direction}})

        # Fetch data from Elasticsearch
        # HIGH FIX #12: Cap export at 10k records to prevent OOM
        MAX_EXPORT_SIZE = 10000
        response = es.search(index=index_name, body=query, size=MAX_EXPORT_SIZE)
        hits = response.get('hits', {}).get('hits', [])
        data = [hit["_source"] for hit in hits]

        # Process Data for Readability
        for record in data:
            timestamp = record.get("@timestamp", "")
            if timestamp:
                try:
                    utc_time = datetime.strptime(timestamp, "%Y-%m-%dT%H:%M:%S.%fZ")
                    ist_time = utc_time + timedelta(hours=5, minutes=30)
                    record["@timestamp"] = ist_time.strftime("%Y-%m-%d %I:%M:%S %p")
                except Exception:
                    record["@timestamp"] = timestamp

            # Handle missing values
            record["UserName"] = record.get("Userdetails", {}).get("UserName", "N/A")
            record["BrokerId"] = record.get("Userdetails", {}).get("BrokerId", "N/A")
            record["LastLoginIp"] = record.get("Userdetails", {}).get("LastLoginIp", "N/A")
            record["LastLoginMac"] = record.get("Userdetails", {}).get("LastLoginMac", "N/A")

        # MEDIUM FIX #21: Streaming Excel with write_only mode — reduced memory footprint
        workbook = Workbook(write_only=True)
        sheet = workbook.create_sheet(title="Filtered User Data")

        headers = ["UserId", "UserName", "BrokerId", "@timestamp", "ReqStatus", "AccessType", "LastLoginIp", "LastLoginMac"]
        sheet.append(headers)

        for record in data:
            sheet.append([
                record.get("UserId", ""),
                record.get("UserName", ""),
                record.get("BrokerId", ""),
                record.get("@timestamp", ""),
                record.get("ReqStatus", ""),
                record.get("AccessType", ""),
                record.get("LastLoginIp", ""),
                record.get("LastLoginMac", "")
            ])

        output = BytesIO()
        workbook.save(output)
        output.seek(0)

        response = HttpResponse(
            output.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename="user_data.xlsx"'
        return response

    except Exception as e:
        return JsonResponse({'error': str(e)}, status=500)


def export_to_pdf(request):
    if request.method != 'POST':
        return JsonResponse({'error': 'Invalid HTTP method. Use POST.'}, status=405)

    try:
        # Parse JSON body
        data = json.loads(request.POST['req'])
        filters = data.get('filters', {})
        sorting = data.get('sorting', [])
        elastic_host = data.get("elastic_host", settings.ELASTIC_HOST)
        elastic_port = data.get("elastic_port", settings.ELASTIC_PORT)

        start_time = data.get("start_time", "none")
        end_time = data.get("end_time", "none")

        # Configure Elasticsearch
        es = Elasticsearch([{'host': elastic_host, 'port': int(elastic_port), 'scheme': 'http'}], 
                           http_auth=(settings.ELASTIC_USER, settings.ELASTIC_PASS))

        index_name = 'noren-login-history'

        # Build Elasticsearch query for counting the matching records
        count_query = {
            "query": {
                "bool": {
                    "must": []
                }
            }
        }

        # Apply date filters if present
        if start_time and end_time:
            count_query["query"]["bool"]["must"].append({
                "range": {
                    "@timestamp": {
                        "gte": start_time,
                        "lte": end_time,
                        "format": "strict_date_optional_time"
                    }
                }
            })

        # Apply filters for the columns
        for col, value in filters.items():
            if value:
                if col in ["UserId", "UserName", "BrokerId", "LastLoginIp", "LastLoginMac"]:
                    count_query["query"]["bool"]["must"].append({
                        "wildcard": {
                            f"Userdetails.{col}.keyword": {
                                "value": f"*{value}*",
                                "case_insensitive": True
                            }
                        }
                    })
                elif col == "@timestamp":
                    count_query["query"]["bool"]["must"].append({
                        "range": {
                            "@timestamp": {
                                "gte": start_time,
                                "lte": end_time,
                                "format": "strict_date_optional_time"
                            }
                        }
                    })

        # Fetch the total number of matching records (total hits)
        count_response = es.count(index=index_name, body=count_query)
        total_records = count_response["count"]  # Get the total count of matching records

        # Build the Elasticsearch query to fetch the actual records (no pagination here)
        query = {
            "_source": [
                "UserId",
                "Userdetails.UserName",
                "Userdetails.BrokerId",
                "@timestamp",
                "ReqStatus",
                "AccessType",
                "Userdetails.LastLoginIp",
                "Userdetails.LastLoginMac"
            ],
            "size": min(total_records, 10000),  # HIGH FIX #12: Cap at 10k to prevent OOM
            "query": {
                "bool": {
                    "must": []
                }
            },
        }

        # Apply date range filter to the main query
        if start_time and end_time:
            query["query"]["bool"]["must"].append({
                "range": {
                    "@timestamp": {
                        "gte": start_time,
                        "lte": end_time,
                        "format": "strict_date_optional_time"
                    }
                }
            })

        # Apply filters for the columns
        for col, value in filters.items():
            if value:
                if col in ["UserId", "UserName", "BrokerId", "LastLoginIp", "LastLoginMac"]:
                    query["query"]["bool"]["must"].append({
                        "wildcard": {
                            f"Userdetails.{col}.keyword": {
                                "value": f"*{value}*",
                                "case_insensitive": True
                            }
                        }
                    })
                elif col == "@timestamp":
                    query["query"]["bool"]["must"].append({
                        "range": {
                            "@timestamp": {
                                "gte": start_time,
                                "lte": end_time,
                                "format": "strict_date_optional_time"
                            }
                        }
                    })

        # Apply sorting
        for sort in sorting:
            column = sort.get('column', '').replace('.', '_')
            direction = sort.get('dir', 'asc')
            if column in ["UserId", "UserName", "BrokerId", "LastLoginIp", "LastLoginMac"]:
                column = f"Userdetails.{column}.keyword"
            query["sort"] = [{column: {"order": direction}}]

        # Fetch data from Elasticsearch
        response = es.search(index=index_name, body=query)
        hits = response.get('hits', {}).get('hits', [])
        data = [hit["_source"] for hit in hits]

        # Process Data for Readability (convert timestamps to human-readable format)
        for record in data:
            timestamp = record.get("@timestamp", "")
            if timestamp:
                try:
                    utc_time = datetime.strptime(timestamp, "%Y-%m-%dT%H:%M:%S.%fZ")
                    ist_time = utc_time + timedelta(hours=5, minutes=30)
                    record["@timestamp"] = ist_time.strftime("%Y-%m-%d %I:%M:%S %p")
                except Exception as e:
                    record["@timestamp"] = timestamp

            # Handle missing values
            record["UserName"] = record.get("Userdetails", {}).get("UserName", "N/A")
            record["BrokerId"] = record.get("Userdetails", {}).get("BrokerId", "N/A")
            record["LastLoginIp"] = record.get("Userdetails", {}).get("LastLoginIp", "N/A")
            record["LastLoginMac"] = record.get("Userdetails", {}).get("LastLoginMac", "N/A")

        # Generate PDF response
        pdf_response = HttpResponse(content_type='application/pdf')
        pdf_response['Content-Disposition'] = 'attachment; filename="user_data.pdf"'

        pdf = canvas.Canvas(pdf_response, pagesize=landscape(letter))
        pdf.setFont("Helvetica-Bold", 10)  # Reduced font size for header
        pdf.drawString(50, 550, "Filtered User Data")
        pdf.setFont("Helvetica", 8)  # Reduced font size for table content

        # Define headers for the table
        headers = ["UserId", "UserName", "BrokerId", "@timestamp", "ReqStatus", "AccessType", "LastLoginIp", "LastLoginMac"]

        # Prepare the table data with records
        table_data = [headers] + [[
            str(record.get("UserId", "")),
            str(record.get("UserName", "")),
            str(record.get("BrokerId", "")),
            str(record.get("@timestamp", "")),
            str(record.get("ReqStatus", "")[:20]),  # Limit to 20 characters
            str(record.get("AccessType", "")),
            str(record.get("LastLoginIp", "")),
            str(record.get("LastLoginMac", ""))
        ] for record in data]

        # Dynamically calculate column widths based on content length
        max_page_width = 750  # Approximate available width in landscape(letter)
        col_widths = []

        for i in range(len(headers)):
            max_content_length = max(len(str(row[i])) for row in table_data)
            col_width = max_content_length * 6  # Adjust multiplier for font size
            col_widths.append(col_width)

        # Scale column widths if they exceed the page width
        total_table_width = sum(col_widths)
        if total_table_width > max_page_width:
            scale_factor = max_page_width / total_table_width
            col_widths = [int(w * scale_factor) for w in col_widths]

        # Table rendering logic
        y = 550  # Initial Y-position for the table
        row_height = 15  # Standard row height
        for i, row in enumerate(table_data):
            x = 50  # Starting X-position
            bg_color = colors.whitesmoke if i % 2 == 0 else colors.beige
            pdf.setFillColor(bg_color)
            pdf.rect(x, y - row_height, sum(col_widths), row_height, stroke=0, fill=1)

            for j, cell in enumerate(row):
                cell_width = col_widths[j]
                pdf.setFillColor(colors.black)
                cell_text = str(cell).strip()
                pdf.drawString(x + 2, y - row_height + 3, cell_text)  # Adjust position for better alignment
                x += cell_width

            y -= row_height
            if y < 50:  # Check if content exceeds page height
                pdf.showPage()
                pdf.setFont("Helvetica", 8)  # Use smaller font for subsequent pages
                y = 550

        pdf.save()

        return pdf_response

    except Exception as e:
        return JsonResponse({'error': str(e)}, status=500)
